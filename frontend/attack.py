from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_SECTION

doc = Document()

# ─────────────────────────────────────────
# PAGE SETUP
# ─────────────────────────────────────────
section = doc.sections[0]
section.page_width    = Inches(8.27)
section.page_height   = Inches(11.69)
section.left_margin   = Cm(2.54)
section.right_margin  = Cm(2.54)
section.top_margin    = Cm(2.54)
section.bottom_margin = Cm(2.54)

# ─────────────────────────────────────────
# MINIMAL PROFESSIONAL COLOUR PALETTE
# Only black, one navy accent, and grays
# ─────────────────────────────────────────
BLACK   = RGBColor(0x1A, 0x1A, 0x1A)
NAVY    = RGBColor(0x1F, 0x3A, 0x5F)   # single accent colour only
DGRAY   = RGBColor(0x4D, 0x4D, 0x4D)
LGRAY   = RGBColor(0xE8, 0xE8, 0xE8)
HGRAY   = RGBColor(0xF5, 0xF5, 0xF5)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
RULE    = RGBColor(0xB0, 0xB0, 0xB0)

# ─────────────────────────────────────────
# DEFAULT STYLE
# ─────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.font.color.rgb = BLACK
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.15

# ─────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────
def set_cell_bg(cell, rgb: RGBColor):
    hex_color = '{:02X}{:02X}{:02X}'.format(rgb.red, rgb.green, rgb.blue)
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top',top),('bottom',bottom),('left',left),('right',right)):
        node = OxmlElement(f'w:{tag}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)

def add_para(text='', bold=False, italic=False, size=11, color=BLACK,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=8,
             font='Times New Roman', line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return p

def add_heading(text, level=1):
    if level == 1:
        p = add_para(text, bold=True, size=14, color=BLACK,
                      space_before=18, space_after=8)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '2')
        bottom.set(qn('w:color'), '{:02X}{:02X}{:02X}'.format(BLACK.red, BLACK.green, BLACK.blue))
        pBdr.append(bottom)
        pPr.append(pBdr)
    elif level == 2:
        add_para(text, bold=True, size=12, color=NAVY,
                  space_before=12, space_after=6)
    elif level == 3:
        add_para(text, bold=True, italic=True, size=11, color=DGRAY,
                  space_before=8, space_after=4)
    return None

def add_bullet(text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    return p

def make_table(headers, rows, col_widths=None, header_bg=NAVY, header_fg=WHITE):
    num_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # header row - single accent colour, no rainbow
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_bg(cell, header_bg)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(2)
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = header_fg
        run.font.name = 'Calibri'

    # data rows - plain white / very light gray, NOT alternating bright colours
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        bg = WHITE if r_idx % 2 == 0 else HGRAY
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_bg(cell, bg)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            run = p.add_run(str(cell_text))
            run.font.size = Pt(10)
            run.font.color.rgb = BLACK
            run.font.name = 'Calibri'

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table

def add_hr(thickness='6', color=RULE):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), thickness)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '{:02X}{:02X}{:02X}'.format(color.red, color.green, color.blue))
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    return p

def page_break():
    doc.add_page_break()


# ═══════════════════════════════════════════════════════════════
# COVER PAGE - clean, formal, single accent colour only
# ═══════════════════════════════════════════════════════════════

# top thin rule
add_hr(thickness='12', color=BLACK)

doc.add_paragraph().paragraph_format.space_after = Pt(60)

add_para('FEASIBILITY REPORT', bold=True, size=26, color=BLACK,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=10,
          font='Times New Roman')

add_para('AmazonShop.lk E-Commerce Web Application',
          bold=False, size=15, color=DGRAY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)

add_para('Prepared in accordance with PRINCE2 Project Management Methodology',
          bold=False, italic=True, size=10, color=DGRAY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=60)

# thin centred rule under subtitle
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('-')
run.font.size = Pt(14)
run.font.color.rgb = RULE
p.paragraph_format.space_after = Pt(60)

doc.add_paragraph().paragraph_format.space_after = Pt(40)

# cover details block - centred, minimal
cover_fields = [
    ('Document Reference', 'AS-FEAS-2026-01'),
    ('Version',            '1.0'),
    ('Prepared By',        'Simthass Mohammed - Project Manager'),
    ('Prepared For',       'Afrath Mohammed - Managing Director, AmazonShop.lk'),
    ('Date',               '17 June 2026'),
]

cov_table = doc.add_table(rows=len(cover_fields), cols=2)
cov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (label, value) in enumerate(cover_fields):
    row = cov_table.rows[i]
    row.cells[0].width = Cm(6)
    row.cells[1].width = Cm(8)
    lp = row.cells[0].paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    lr = lp.add_run(label + ' :')
    lr.bold = True
    lr.font.size = Pt(10.5)
    lr.font.color.rgb = DGRAY
    lr.font.name = 'Calibri'
    vp = row.cells[1].paragraphs[0]
    vp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    vr = vp.add_run('  ' + value)
    vr.font.size = Pt(10.5)
    vr.font.color.rgb = BLACK
    vr.font.name = 'Calibri'
    # remove borders for clean look
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'), 'nil')
            borders.append(el)
        tcPr.append(borders)

doc.add_paragraph().paragraph_format.space_after = Pt(100)

# bottom institutional line
add_para('University of Bedfordshire (SLIIT City University)',
          size=10, color=DGRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para('Agile Project Management - CIS047-3',
          size=10, color=DGRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

# bottom thin rule
add_hr(thickness='12', color=BLACK)

page_break()


# ═══════════════════════════════════════════════════════════════
# DOCUMENT CONTROL
# ═══════════════════════════════════════════════════════════════
add_heading('Document Control', 1)

add_heading('Revision History', 2)
rev_rows = [
    ('0.1', '15 June 2026', 'Simthass Mohammed', 'Initial draft'),
    ('1.0', '17 June 2026', 'Simthass Mohammed', 'Final version for approval'),
]
make_table(['Version', 'Date', 'Author', 'Description'], rev_rows,
           col_widths=[2.2, 3.8, 4.5, 6.5])

add_heading('Approval', 2)
appr_rows = [
    ('Afrath Mohammed',   'Client / Managing Director', 'Approved to Proceed', '17 June 2026'),
    ('Simthass Mohammed', 'Project Manager',            'Recommended',          '17 June 2026'),
]
make_table(['Name', 'Role', 'Decision', 'Date'], appr_rows,
           col_widths=[4.5, 5, 4.5, 3])

add_heading('Distribution List', 2)
dist_rows = [
    ('Afrath Mohammed',   'Managing Director, AmazonShop.lk'),
    ('Simthass Mohammed', 'Project Manager'),
    ('Project Team',      'Start-up, Quality, Risk, Scheduling Managers'),
]
make_table(['Name', 'Role'], dist_rows, col_widths=[6, 11])

page_break()


# ═══════════════════════════════════════════════════════════════
# TABLE OF CONTENTS (static)
# ═══════════════════════════════════════════════════════════════
add_heading('Table of Contents', 1)

toc_items = [
    ('1.', 'Introduction'),
    ('1.1', 'Purpose'),
    ('1.2', 'Background'),
    ('2.', 'Technical Feasibility'),
    ('3.', 'Economic Feasibility'),
    ('4.', 'Legal Feasibility'),
    ('5.', 'Operational Feasibility'),
    ('6.', 'Scheduling Feasibility'),
    ('7.', 'Overall Feasibility Summary'),
    ('8.', 'Recommendation'),
    ('9.', 'Approval'),
]
for num, title in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    is_main = num.replace('.', '').isdigit() and '.' not in num[:-1]
    p.paragraph_format.left_indent = Cm(0) if is_main else Cm(0.8)
    r1 = p.add_run(f'{num}  ')
    r1.bold = is_main
    r1.font.size = Pt(11)
    r1.font.color.rgb = BLACK
    r1.font.name = 'Times New Roman'
    r2 = p.add_run(title)
    r2.bold = is_main
    r2.font.size = Pt(11)
    r2.font.color.rgb = BLACK
    r2.font.name = 'Times New Roman'

page_break()


# ═══════════════════════════════════════════════════════════════
# 1. INTRODUCTION
# ═══════════════════════════════════════════════════════════════
add_heading('1. Introduction', 1)

add_heading('1.1 Purpose', 2)
add_para(
    'This Feasibility Report has been prepared to assess whether the proposed AmazonShop.lk '
    'E-Commerce Web Application is technically achievable, economically viable, legally compliant, '
    'operationally practical, and deliverable within the available project timeframe. The findings '
    'of this report support the continued business justification of the project as required by the '
    'PRINCE2 methodology and inform the decision to proceed to the detailed requirements and '
    'development stages.'
)

add_heading('1.2 Background', 2)
add_para(
    'AmazonShop.lk currently operates exclusively as a physical retail store in Colombo, Sri Lanka, '
    'selling mobile phones, mobile accessories, and desktop devices. Following a formal system '
    'request from the Managing Director, the project team has been tasked with evaluating and '
    'subsequently developing a full e-commerce platform to support online sales.'
)
add_para(
    'Before proceeding to full-scale development, this report evaluates whether the proposed '
    'solution is feasible across five key dimensions: Technical, Economic, Legal, Operational, '
    'and Scheduling, collectively referred to as the TELOS framework.'
)

page_break()


# ═══════════════════════════════════════════════════════════════
# 2. TECHNICAL FEASIBILITY
# ═══════════════════════════════════════════════════════════════
add_heading('2. Technical Feasibility', 1)

add_heading('2.1 Overview', 2)
add_para(
    'Technical feasibility assesses whether the project team possesses the necessary technology, '
    'tools, and technical expertise to successfully build the proposed system.'
)

add_heading('2.2 Assessment', 2)
tech_rows = [
    ('Development Team Skills', 'The team has working knowledge of React, Node.js, Express, and MongoDB through prior coursework and personal projects'),
    ('Technology Maturity', 'React, Node.js, and MongoDB are mature, well-documented, widely used technologies with extensive community support'),
    ('Third-Party Services', 'Required services (Cloudinary, Stripe/PayHere, Google Gemini API, SendGrid) all offer free tiers suitable for development and demonstration'),
    ('Hosting Availability', 'Free and low-cost hosting options (Vercel, Render, MongoDB Atlas) are sufficient for this project\u2019s scale'),
    ('Infrastructure Requirements', 'No specialised hardware is required; standard laptops with internet connectivity are sufficient for all team members'),
    ('AI Chatbot Complexity', 'The Google Gemini API provides an accessible interface for natural language processing, making chatbot implementation realistic within the timeframe'),
]
make_table(['Factor', 'Assessment'], tech_rows, col_widths=[5, 12])

add_heading('2.3 Technical Risks Identified', 2)
for item in [
    'Payment gateway integration is the most technically complex component and carries the highest risk of delay',
    'Team members have varying levels of prior experience with full-stack development',
    'Free-tier API rate limits (Gemini, Cloudinary) may need monitoring during testing',
]:
    add_bullet(item)

add_heading('2.4 Conclusion', 2)
add_para(
    'The project is technically feasible. All required technologies are mature, accessible, and '
    'within the team\u2019s collective skill set. Identified risks, primarily around payment integration, '
    'are manageable through early-stage testing and have been recorded in the Risk Management Plan.'
)

page_break()


# ═══════════════════════════════════════════════════════════════
# 3. ECONOMIC FEASIBILITY
# ═══════════════════════════════════════════════════════════════
add_heading('3. Economic Feasibility', 1)

add_heading('3.1 Overview', 2)
add_para(
    'Economic feasibility assesses whether the financial costs of the project are justified by the '
    'expected benefits, and whether the project represents a worthwhile investment.'
)

add_heading('3.2 Cost Analysis', 2)
cost_rows = [
    ('Development Labour', 'LKR 0', 'Performed by student project team as part of coursework'),
    ('Hosting - Frontend (Vercel)', 'LKR 0', 'Free tier sufficient for project scale'),
    ('Hosting - Backend (Render)', 'LKR 0', 'Free tier sufficient for development and demonstration'),
    ('Database (MongoDB Atlas)', 'LKR 0', 'M0 free cluster sufficient for project scale'),
    ('Image Storage (Cloudinary)', 'LKR 0', 'Free tier covers required storage and bandwidth'),
    ('Payment Gateway (Stripe/PayHere)', 'Transaction fees only (2.5\u20133.5% per transaction)', 'No upfront cost; fees apply only on live transactions'),
    ('AI Chatbot (Gemini API)', 'LKR 0', 'Free tier sufficient for demonstration volume'),
    ('Domain Name (future production)', 'Approx. LKR 3,000\u20135,000 / year', 'Optional for academic submission; required for real-world launch'),
]
make_table(['Cost Category', 'Estimated Cost', 'Notes'], cost_rows, col_widths=[5, 4.5, 7.5])
add_para('Total Estimated Cost (Academic Phase): Approximately LKR 0 - all core infrastructure uses free tiers.',
          bold=True, size=10.5, color=BLACK)

add_heading('3.3 Expected Benefits', 2)
ben_rows = [
    ('New Revenue Channel', 'Enables 24/7 online sales beyond physical store hours and location'),
    ('Market Expansion', 'Reaches customers islandwide rather than only Colombo walk-in traffic'),
    ('Operational Efficiency', 'Reduces manual order-taking and inventory tracking errors'),
    ('Competitive Positioning', 'Matches or exceeds competitors who may already have an online presence'),
    ('Data Insights', 'Enables data-driven decisions through sales and customer analytics'),
    ('Estimated Revenue Impact', 'Based on industry benchmarks for SME e-commerce adoption in Sri Lanka, a realistic estimate of 20\u201340% increase in monthly revenue within 6 months of launch is achievable'),
]
make_table(['Benefit', 'Description'], ben_rows, col_widths=[5, 12])

add_heading('3.4 Cost-Benefit Summary', 2)
add_para(
    'Given that the academic phase of this project carries effectively zero direct financial cost, '
    'while the potential business benefit to AmazonShop.lk is a meaningful new revenue channel, '
    'the cost-benefit ratio is highly favourable. For future production scaling beyond the academic '
    'project, ongoing costs would include paid hosting tiers, a domain name, and payment gateway '
    'transaction fees, all of which are standard, modest costs for an SME e-commerce operation in '
    'Sri Lanka.'
)

add_heading('3.5 Conclusion', 2)
add_para(
    'The project is economically feasible. The academic development phase incurs negligible direct '
    'cost, while the long-term business benefit to AmazonShop.lk justifies the investment of time '
    'and effort.'
)

page_break()


# ═══════════════════════════════════════════════════════════════
# 4. LEGAL FEASIBILITY
# ═══════════════════════════════════════════════════════════════
add_heading('4. Legal Feasibility', 1)

add_heading('4.1 Overview', 2)
add_para(
    'Legal feasibility assesses whether the proposed system can be developed and operated in '
    'compliance with relevant laws, regulations, and industry standards.'
)

add_heading('4.2 Assessment', 2)
legal_rows = [
    ('Data Protection', 'Customer personal data must be handled responsibly', 'Secure password hashing, HTTPS encryption, minimal PII storage, and a published Privacy Policy'),
    ('PCI DSS Compliance', 'Direct handling of card data carries significant compliance burden', 'Card transactions handled entirely by the PCI-compliant third-party gateway (Stripe/PayHere)'),
    ('Consumer Protection', 'Customers must be informed of pricing, returns, and order terms', 'Terms and Conditions and Return Policy pages published and linked from checkout and footer'),
    ('Intellectual Property', 'Product images, logos, and content must not infringe third-party rights', 'Images supplied by the client or used with appropriate manufacturer rights'),
    ('Sri Lankan E-Commerce Regulations', 'No dedicated e-commerce law, but general consumer protection and Computer Crimes Act apply', 'System follows OWASP security best practices to mitigate liability under the Computer Crimes Act No. 24 of 2007'),
    ('Academic Integrity', 'Project must not replicate any individual member\u2019s prior UG project work', 'Confirmed as an original group project developed specifically for CIS047-3'),
]
make_table(['Legal Area', 'Consideration', 'Compliance Approach'], legal_rows, col_widths=[4, 5.5, 7.5])

add_heading('4.3 Conclusion', 2)
add_para(
    'The project is legally feasible. By using a third-party PCI-compliant payment gateway and '
    'following standard data protection and security practices, the system avoids the most '
    'significant legal compliance burdens while remaining consistent with Sri Lankan consumer '
    'protection expectations.'
)

page_break()


# ═══════════════════════════════════════════════════════════════
# 5. OPERATIONAL FEASIBILITY
# ═══════════════════════════════════════════════════════════════
add_heading('5. Operational Feasibility', 1)

add_heading('5.1 Overview', 2)
add_para(
    'Operational feasibility assesses whether the proposed system will function effectively within '
    'the day-to-day business operations of AmazonShop.lk once deployed, and whether staff and '
    'customers will be able to use it effectively.'
)

add_heading('5.2 Assessment', 2)
op_rows = [
    ('Staff Technical Capability', 'Shop staff have basic computer literacy; the admin panel requires no programming knowledge and is operable after minimal (30-minute) familiarisation'),
    ('Customer Adoption', 'Sri Lanka has high smartphone and internet penetration; target customers (mobile/tech buyers) are likely to be comfortable with online shopping'),
    ('Order Fulfilment Process', 'The existing physical store provides inventory and stock-handling infrastructure that can be adapted to fulfil online orders'),
    ('Delivery Logistics', 'The client will need to establish or partner with a delivery service; this is a business process decision outside the technical system but necessary for operational success'),
    ('Customer Support', 'The AI chatbot reduces basic support burden; complex queries are routed to the Contact page for human follow-up'),
    ('Change Management', 'Staff will require basic training on the admin dashboard, provided as part of project handover documentation'),
]
make_table(['Factor', 'Assessment'], op_rows, col_widths=[5, 12])

add_heading('5.3 Organisational Impact', 2)
add_para(
    'The introduction of the e-commerce platform represents a moderate but manageable change to '
    'AmazonShop.lk\u2019s operations. Staff will need to incorporate order processing and inventory '
    'synchronisation between the physical store and online platform into their daily routine. This '
    'is a common and well-understood operational shift for small retail businesses transitioning to '
    'omnichannel sales.'
)

add_heading('5.4 Conclusion', 2)
add_para(
    'The project is operationally feasible. The system has been designed specifically with '
    'non-technical staff usability in mind, and the target customer base is well-suited to online '
    'shopping adoption. The client should plan delivery logistics as a parallel business decision.'
)

page_break()


# ═══════════════════════════════════════════════════════════════
# 6. SCHEDULING FEASIBILITY
# ═══════════════════════════════════════════════════════════════
add_heading('6. Scheduling Feasibility', 1)

add_heading('6.1 Overview', 2)
add_para(
    'Scheduling feasibility assesses whether the project can realistically be completed within the '
    'available timeframe given the scope of work and available resources.'
)

add_heading('6.2 Assessment', 2)
sched_rows = [
    ('Available Timeframe', '10 weeks, as defined by the CIS047-3 module structure'),
    ('Team Size', '5 members, each with a clearly defined ownership area'),
    ('Development Approach', 'Agile Scrum with four 2-week sprints, allowing incremental delivery and early identification of delays'),
    ('Feature Scope', 'Core e-commerce features follow well-established patterns with extensive documentation available, reducing development time'),
    ('Highest Risk Area', 'Payment gateway integration (Sprint 3) is the most time-sensitive component and has buffer time scheduled before the Sprint 3 review'),
    ('Parallel Workstreams', 'Distinct vertical feature ownership allows work to proceed in parallel, maximising use of the 10-week window'),
]
make_table(['Factor', 'Assessment'], sched_rows, col_widths=[5, 12])

add_heading('6.3 Sprint Allocation Summary', 2)
sprint_rows = [
    ('Sprint 1', 'Weeks 1\u20132', 'Project setup, authentication foundation, database schema'),
    ('Sprint 2', 'Weeks 3\u20134', 'Core public pages with real data, cart functionality'),
    ('Sprint 3', 'Weeks 5\u20136', 'Checkout, payment integration, chatbot, admin panel core'),
    ('Sprint 4', 'Weeks 7\u20138', 'Testing, polish, responsive design, demo video'),
    ('Closure',  'Weeks 9\u201310', 'Documentation finalisation and submission'),
]
make_table(['Sprint', 'Weeks', 'Focus'], sprint_rows, col_widths=[3, 3, 11])

add_heading('6.4 Conclusion', 2)
add_para(
    'The project is feasible within the scheduled timeframe. The 10-week window, combined with a '
    'five-member team working on parallel feature ownership and an Agile sprint structure, provides '
    'sufficient time to deliver the defined scope, provided payment integration begins early in '
    'Sprint 3 as planned.'
)

page_break()


# ═══════════════════════════════════════════════════════════════
# 7. OVERALL FEASIBILITY SUMMARY
# ═══════════════════════════════════════════════════════════════
add_heading('7. Overall Feasibility Summary', 1)

summary_rows = [
    ('Technical',    'Feasible', 'Mature technologies, accessible APIs, manageable complexity'),
    ('Economic',     'Feasible', 'Negligible direct cost; favourable cost-benefit ratio'),
    ('Legal',        'Feasible', 'PCI compliance handled by third-party gateway; standard data protection practices sufficient'),
    ('Operational',  'Feasible', 'Designed for non-technical staff use; suitable target customer base'),
    ('Scheduling',   'Feasible', '10-week timeframe sufficient with parallel team structure and Agile sprints'),
]
make_table(['Dimension', 'Rating', 'Summary'], summary_rows, col_widths=[4, 3, 10])

page_break()


# ═══════════════════════════════════════════════════════════════
# 8. RECOMMENDATION
# ═══════════════════════════════════════════════════════════════
add_heading('8. Recommendation', 1)

add_para(
    'Based on the assessment across all five dimensions of the TELOS framework, this report '
    'concludes that the proposed AmazonShop.lk E-Commerce Web Application is feasible and '
    'recommended to proceed to the detailed requirements specification and development stages.'
)

add_para('The project team should pay particular attention to the following:', space_after=4)
for item in [
    'Beginning payment gateway integration early in Sprint 3 to mitigate the highest-risk technical component',
    'Ensuring the client provides product data and brand assets promptly to avoid content-related delays',
    'Maintaining the Agile sprint structure to allow early detection of any scheduling slippage',
]:
    add_bullet(item)

page_break()


# ═══════════════════════════════════════════════════════════════
# 9. APPROVAL
# ═══════════════════════════════════════════════════════════════
add_heading('9. Approval', 1)

add_para(
    'This Feasibility Report has been reviewed and approved by the undersigned, confirming '
    'agreement to proceed with the AmazonShop.lk E-Commerce Web Application project.'
)

final_appr_rows = [
    ('Afrath Mohammed',   'Client / Managing Director', 'Approved to Proceed', '____________', '17 June 2026'),
    ('Simthass Mohammed', 'Project Manager',            'Recommended',          '____________', '17 June 2026'),
]
make_table(['Name', 'Role', 'Decision', 'Signature', 'Date'], final_appr_rows,
           col_widths=[4, 4.5, 4, 3, 2.5])

doc.add_paragraph()
add_hr(thickness='6', color=RULE)
add_para('End of Feasibility Report  |  Document Ref: AS-FEAS-2026-01  |  Version 1.0  |  AmazonShop.lk  |  17 June 2026',
          size=9, color=DGRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4)


# ═══════════════════════════════════════════════════════════════
# HEADER & FOOTER - minimal, professional
# ═══════════════════════════════════════════════════════════════
def add_header_footer():
    sec = doc.sections[0]

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr1 = hp.add_run('AmazonShop.lk  |  Feasibility Report  |  AS-FEAS-2026-01')
    hr1.font.size = Pt(8)
    hr1.font.color.rgb = DGRAY
    hr1.font.name = 'Calibri'
    hp.paragraph_format.space_after = Pt(0)

    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '{:02X}{:02X}{:02X}'.format(RULE.red, RULE.green, RULE.blue))
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_left = fp.add_run('Confidential - AmazonShop.lk Project Documentation     ')
    run_left.font.size = Pt(8)
    run_left.font.color.rgb = DGRAY
    run_left.font.name = 'Calibri'

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run_page = fp.add_run()
    run_page.font.size = Pt(8)
    run_page.font.color.rgb = DGRAY
    run_page._r.append(fldChar1)
    run_page._r.append(instrText)
    run_page._r.append(fldChar2)

add_header_footer()

# ─────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────
output_path = 'AmazonShop_Feasibility_Report_v1.0.docx'
doc.save(output_path)
print(f'Document generated successfully: {output_path}')
